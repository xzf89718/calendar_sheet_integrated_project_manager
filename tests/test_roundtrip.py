# -*- coding: utf-8 -*-
"""导入导出往返一致性 + 单元格转换规则。

往返 = 导出 query_projects 的行 -> 新库导入 -> 行集（业务字段）相等。
导出/导入与查询同管线，这是"导出所见 = 筛选结果"合同的回归锁。
"""
from datetime import datetime

from openpyxl import load_workbook

from db import normalize_date
from exporter import export_excel, export_with_template
from importer import _cell_to_field, import_file


def _row_dicts(rows, fields):
    return sorted(
        tuple(str(r.get(f) or "") for f in fields) for r in rows)


def test_export_import_roundtrip(seeded, dbs, tmp_path, cfg):
    rows, _ = seeded.query_projects(include_closed=True, sort="id")
    path = tmp_path / "out.xlsx"
    n = export_excel(rows, str(path))
    assert n == 5

    db2_path = tmp_path / "roundtrip.db"
    import db as db_mod
    db2 = db_mod.Database(str(db2_path))
    try:
        created, skipped, unmapped = import_file(db2, str(path))
        assert created == 5 and skipped == [] and unmapped == []
        rows2, total2 = db2.query_projects(include_closed=True, sort="id")
        assert total2 == 5
        fields = ["manufacturer", "customer", "product", "amount",
                  "qty", "start_date", "next_followup", "close_date",
                  "stage", "status", "note"]
        assert _row_dicts(rows, fields) == _row_dicts(rows2, fields)
    finally:
        db2.close()


def test_roundtrip_dates_survive(seeded, tmp_path):
    rows, _ = seeded.query_projects(include_closed=True)
    path = tmp_path / "d.xlsx"
    export_excel(rows, str(path))
    db2_path = tmp_path / "rt2.db"
    import db as db_mod
    db2 = db_mod.Database(str(db2_path))
    try:
        import_file(db2, str(path))
        rows2, _ = db2.query_projects(include_closed=True, sort="id")
        assert [r["start_date"] for r in rows2] == \
               [r["start_date"] for r in sorted(rows, key=lambda r: r["id"])]
    finally:
        db2.close()


def test_export_template_const_column(seeded, cfg, tmp_path):
    rows, _ = seeded.query_projects(include_closed=True)
    template = {
        "name": "汇报",
        "columns": [("厂家", "manufacturer", None),
                    ("拥有者", None, "张三"),
                    ("备注", "note", None)],
        "defaults": {},
    }
    path = tmp_path / "t.xlsx"
    export_with_template(rows, str(path), template)
    ws = load_workbook(str(path)).active
    data = list(ws.iter_rows(values_only=True))
    assert data[0] == ("厂家", "拥有者", "备注")
    assert all(r[1] == "张三" for r in data[1:])
    assert len(data) == 6                        # 表头 + 5 行


def test_import_csv_with_unmapped_header(seeded, dbs, tmp_path):
    p = tmp_path / "in.csv"
    p.write_text("厂家,客户,产品,不存在的列\n丁,戊,己,垃圾\n",
                 encoding="utf-8-sig")
    created, skipped, unmapped = import_file(dbs, str(p))
    assert created == 1
    assert unmapped == ["不存在的列"]
    rows, total = dbs.query_projects(keywords=["丁"],
                                     search_fields=["manufacturer"])
    assert total == 1


def test_import_skips_blank_and_repeated_header(seeded, dbs, tmp_path):
    p = tmp_path / "in.csv"
    p.write_text("厂家,客户,产品\n,,,\n厂家,客户,产品\n庚,辛,壬\n",
                 encoding="utf-8-sig")
    created, skipped, _u = import_file(dbs, str(p))
    # ",,,"行在读入时被整行过滤（importer 只留非空行），
    # 重复表头行被识别跳过 -> skipped = [2]
    assert created == 1 and skipped == [2]


# ---------- _cell_to_field 转换规则 ----------

def test_cell_datetime(cfg):
    assert _cell_to_field("start_date", datetime(2026, 3, 1)) == "2026-03-01"
    assert _cell_to_field("note", datetime(2026, 3, 1)) == "2026-03-01 00:00:00"


def test_cell_excel_serial_date(cfg):
    # Excel 序列号 46082 = 2026-03-01（1899-12-30 起算）
    assert _cell_to_field("start_date", 46082) == "2026-03-01"


def test_cell_numeric_strings(cfg):
    assert _cell_to_field("qty", "25%") == 25
    assert _cell_to_field("amount", "3.14") == 3.14
    assert _cell_to_field("qty", "abc") is None


def test_cell_empty_variants(cfg):
    assert _cell_to_field("note", "") is None
    assert _cell_to_field("note", "None") is None
    assert _cell_to_field("note", None) is None
    assert _cell_to_field("note", "  保留  ") == "保留"


def test_cell_date_string_normalized(cfg):
    assert _cell_to_field("start_date", "2026年3月1日") == "2026-03-01"
    assert _cell_to_field("start_date", "胡说") is None


def test_normalize_date_is_importer_dependency():
    # normalize_date 是导入管线的根基，独立锁一下行为
    assert normalize_date("２０２６／１２／３１") == "2026-12-31"
    assert normalize_date("12/31/2026") is None     # 非 ISO 顺序不认
