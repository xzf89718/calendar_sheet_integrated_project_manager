# -*- coding: utf-8 -*-
"""周报清洗器（子模块）：预览原始周报 Excel，按日期匹配规则清洗后导出。

用法：python weekreport_clean/clean_gui.py
规则：第一列（修改日期）规范化后出现在"跟进记录"列文本中的行保留。
配置：跟进记录列名优先取根目录 fields.json 周报模板的对应标签。
"""
import os
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from openpyxl import Workbook, load_workbook      # noqa: E402
from config import CFG                             # noqa: E402
import weekreport_clean.clean_core as core         # noqa: E402


class CleanApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("周报清洗器 v1.0.0")
        self.geometry("1100x680")
        self.headers = []
        self.rows = []
        self.kept = []
        self.note_idx = None
        self._build()

    def _build(self):
        pad = {"padx": 8, "pady": 4}
        top = tk.Frame(self)
        top.pack(fill="x", **pad)
        tk.Label(top, text="输入文件：").pack(side="left")
        self.path_var = tk.StringVar()
        tk.Entry(top, textvariable=self.path_var).pack(
            side="left", fill="x", expand=True, padx=4)
        tk.Button(top, text="浏览", command=self._pick).pack(side="left")
        tk.Button(top, text="加载", bg="#1565c0", fg="white",
                  command=self._load).pack(side="left", padx=4)

        up = tk.LabelFrame(self, text="原始内容", padx=6, pady=6)
        up.pack(fill="both", expand=True, **pad)
        self.tree_up = self._make_tree(up)

        mid = tk.Frame(self)
        mid.pack(fill="x", **pad)
        tk.Button(mid, text="清洗预览 →", bg="#2e7d32", fg="white",
                  command=self._clean).pack(side="left", padx=4)
        self.stat = tk.Label(mid, text="先选择并加载周报文件", fg="#666666")
        self.stat.pack(side="left", padx=10)
        tk.Button(mid, text="导出清洗结果", command=self._export
                  ).pack(side="right", padx=4)

        low = tk.LabelFrame(self, text="清洗结果", padx=6, pady=6)
        low.pack(fill="both", expand=True, **pad)
        self.tree_low = self._make_tree(low)

    def _make_tree(self, parent):
        tree = ttk.Treeview(parent, show="headings", height=10)
        ysb = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
        xsb = ttk.Scrollbar(parent, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
        tree.pack(side="left", fill="both", expand=True)
        ysb.pack(side="left", fill="y")
        xsb.pack(fill="x")
        return tree

    # ---------- 动作 ----------
    def _pick(self):
        p = filedialog.askopenfilename(
            parent=self,
            filetypes=[("Excel", "*.xlsx"), ("所有", "*.*")])
        if p:
            self.path_var.set(p)

    def _load(self):
        p = self.path_var.get().strip()
        if not p or not os.path.exists(p):
            messagebox.showwarning("提示", "请先选择存在的文件", parent=self)
            return
        try:
            wb = load_workbook(p, read_only=True)
            ws = wb[wb.sheetnames[0]]
            self.headers, self.rows = core.load_rows(ws)
            wb.close()
        except Exception as e:                          # noqa: BLE001
            messagebox.showerror("读取失败", str(e), parent=self)
            return
        # 跟进记录列：优先周报模板标签，其次包含"跟进"
        prefer = None
        tpl = next((t for t in CFG.templates
                    if t.get("kind") == "weekly"), None)
        if tpl:
            for label, field, _c in tpl["columns"]:
                if field == "followup_note":
                    prefer = label
        self.note_idx = core.find_note_idx(self.headers, prefer)
        self._fill(self.tree_up, self.rows)
        self.tree_low.delete(*self.tree_low.get_children())
        self.stat.config(
            text=f"共 {len(self.rows)} 行 ｜ 跟进记录列: "
                 f"{self.headers[self.note_idx] if self.note_idx is not None else '未找到'}"
                 f" ｜ 点「清洗预览」")

    def _fill(self, tree, rows):
        tree.configure(columns=list(range(len(self.headers))),
                       displaycolumns=list(range(len(self.headers))))
        for i, h in enumerate(self.headers):
            tree.heading(i, text=h)
            tree.column(i, width=180, anchor="w", stretch=False)
        tree.delete(*tree.get_children())
        for r in rows:
            tree.insert("", "end", values=["" if v is None else str(v)
                                           for v in r])

    def _clean(self):
        if not self.rows:
            messagebox.showinfo("提示", "请先加载文件", parent=self)
            return
        if self.note_idx is None:
            messagebox.showwarning(
                "提示", "找不到「跟进记录」列（表头需含'跟进'或匹配"
                       "周报模板标签）", parent=self)
            return
        self.kept, dropped = core.clean_rows(self.headers, self.rows,
                                             self.note_idx)
        self._fill(self.tree_low, self.kept)
        self.stat.config(
            text=f"保留 {len(self.kept)} 行 / 丢弃 {dropped} 行 ｜ "
                 f"规则: 第一列日期出现在跟进记录中才保留")

    def _export(self):
        if not self.kept:
            messagebox.showinfo("提示", "请先清洗预览（结果为空不能导出）",
                                parent=self)
            return
        p = filedialog.asksaveasfilename(
            parent=self, defaultextension=".xlsx",
            initialfile="周报_清洗后.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not p:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "清洗结果"
            ws.append(self.headers)
            for r in self.kept:
                ws.append(list(r))
            wb.save(p)
        except Exception as e:                          # noqa: BLE001
            messagebox.showerror("导出失败", str(e), parent=self)
            return
        messagebox.showinfo("完成", f"已导出：\n{p}", parent=self)


if __name__ == "__main__":
    CleanApp().mainloop()
